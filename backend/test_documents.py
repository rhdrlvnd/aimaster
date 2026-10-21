from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import re
import unittest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from backend.documents import extract


class DocumentTests(unittest.TestCase):
    def test_word_order_and_table(self):
        doc = Document()
        doc.add_paragraph('운영 절차')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '매일 로그 확인'
        doc.add_paragraph('오류 발생 시 연락')
        stream = BytesIO(); doc.save(stream)
        text, warnings = extract('manual.docx', stream.getvalue())
        self.assertLess(text.index('운영 절차'), text.index('매일 로그 확인'))
        self.assertLess(text.index('매일 로그 확인'), text.index('오류 발생 시 연락'))
        self.assertTrue(warnings)

    def test_excel_cells_and_formulas(self):
        book = Workbook(); sheet = book.active; sheet.title = '업무'
        sheet['A1'] = '매일 확인'; sheet['B3'] = '=1+1'
        stream = BytesIO(); book.save(stream); book.close()
        text, warnings = extract('tasks.xlsx', stream.getvalue())
        self.assertIn('시트 업무', text)
        self.assertIn('A1: 매일 확인', text)
        self.assertIn('B3: =1+1', text)
        self.assertTrue(warnings)

    def test_presentation(self):
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = '업무 안내'
        slide.notes_slide.notes_text_frame.text = '장애 시 담당자 연락'
        stream = BytesIO(); deck.save(stream)
        text, _ = extract('guide.pptx', stream.getvalue())
        self.assertIn('슬라이드 1', text)
        self.assertIn('업무 안내', text)
        self.assertIn('장애 시 담당자 연락', text)

    def test_excel_without_dimension_metadata(self):
        book = Workbook(); book.active['A1'] = '매일 로그 점검'
        stream = BytesIO(); book.save(stream); book.close()
        result = BytesIO()
        with ZipFile(stream) as src, ZipFile(result, 'w', ZIP_DEFLATED) as dst:
            for name in src.namelist():
                data = src.read(name)
                if name == 'xl/worksheets/sheet1.xml':
                    data = re.sub(rb'<dimension[^>]*/>', b'', data)
                dst.writestr(name, data)
        self.assertIn('매일 로그 점검', extract('manual.xlsx', result.getvalue())[0])

    def test_invalid_empty_and_old_formats(self):
        for name, data in [('a.docx', b'broken'), ('a.xls', b'legacy'), ('a.txt', b'')]:
            with self.assertRaises(ValueError):
                extract(name, data)
