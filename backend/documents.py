"""Extract Office text without executing macros, formulas, or external links."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pptx import Presentation

SUPPORTED = {'.txt', '.md', '.docx', '.xlsx', '.pptx'}
OFFICE_LIMIT = 10 * 1024 * 1024


def extract(name, data):
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError('DOCX·XLSX·PPTX·TXT·MD를 지원합니다. 구형 DOC·XLS·PPT는 최신 형식으로 저장해 주세요.')
    warnings = []
    parts = []

    def add(label, text):
        if text and text.strip():
            parts.append(f'[{label}]\n{text.strip()}')
            if sum(map(len, parts)) > 45_000:
                raise ValueError('추출한 내용이 45,000자를 초과합니다. 문서를 나누어 등록해 주세요.')

    try:
        if suffix in {'.txt', '.md'}:
            text = data.decode('utf-8-sig')
        else:
            with ZipFile(BytesIO(data)) as archive:
                if len(archive.infolist()) > 5000 or sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError('압축을 푼 문서 크기가 처리 한도를 초과합니다.')
            warnings.append('텍스트와 표를 추출했습니다. 이미지·차트·SmartArt와 원본 배치는 분석에 포함되지 않습니다.')
            if suffix == '.docx':
                doc = Document(BytesIO(data))
                for number, block in enumerate(doc.iter_inner_content(), 1):
                    if isinstance(block, Paragraph):
                        add(f'본문 {number}', block.text)
                    else:
                        for row_number, row in enumerate(block.rows, 1):
                            add(f'표 {number} / 행 {row_number}', ' | '.join(c.text for c in row.cells))
                warnings.append('Word 머리글·바닥글·각주·텍스트 상자는 현재 추출 대상이 아닙니다.')
            elif suffix == '.xlsx':
                book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
                try:
                    for sheet in book:
                        if sheet.max_row is None or sheet.max_column is None:
                            sheet.calculate_dimension(force=True)
                        if sheet.max_row * sheet.max_column > 200_000:
                            raise ValueError('시트의 행·열 범위가 너무 큽니다. 불필요한 빈 영역을 정리해 주세요.')
                        for row_number, row in enumerate(sheet.iter_rows(), 1):
                            values = [f'{c.coordinate}: {c.value}' for c in row if c.value is not None]
                            add(f'시트 {sheet.title} / 행 {row_number}' if values else '', ' | '.join(values))
                finally:
                    book.close()
                warnings.append('숨김 시트·셀도 포함합니다. 수식은 계산하지 않고 수식 문자열로 읽습니다.')
            else:
                deck = Presentation(BytesIO(data))

                def shapes_text(shapes):
                    for shape in shapes:
                        if hasattr(shape, 'shapes'):
                            yield from shapes_text(shape.shapes)
                        if shape.has_text_frame:
                            yield shape.text_frame.text
                        if shape.has_table:
                            for row in shape.table.rows:
                                yield ' | '.join(c.text for c in row.cells)

                for number, slide in enumerate(deck.slides, 1):
                    add(f'슬라이드 {number}', '\n'.join(shapes_text(slide.shapes)))
                    if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                        add(f'슬라이드 {number} / 발표자 노트', slide.notes_slide.notes_text_frame.text)
            text = '\n\n'.join(parts)
    except ValueError:
        raise
    except Exception as error:
        raise ValueError('문서를 읽을 수 없습니다. 손상·암호 설정·파일 형식을 확인해 주세요.') from error
    if not text.strip() and suffix == '.pptx':
        text = '(추출 텍스트 없음 — 슬라이드 이미지 참조)'
    if not text.strip() or '\x00' in text:
        raise ValueError('추출할 텍스트가 없습니다. 이미지로만 된 문서는 현재 읽을 수 없습니다.')
    return text, warnings
